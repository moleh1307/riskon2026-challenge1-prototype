"""Adapted structural-ingestion contracts for the canonical event runtime."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from pydantic import ValidationError

from riskon.config import load_milestone2_config
from riskon.event_runtime.corpus_loader import ingest_event_sections
from riskon.event_runtime.evidence_reasoning import EventEvidenceReasoningRuntime
from riskon.event_runtime.evidence_reasoning_models import (
    ContextAssessment,
    EvidenceAnalysisOutput,
    EvidenceSufficiencyStatus,
    SkepticOutput,
    SupportValidation,
)
from riskon.event_runtime.retrieval import EventHybridRetriever
from riskon.event_structure.acronyms import build_glossary, verify_expansion
from riskon.event_structure.adapter import (
    attach_structural_tables,
    build_structural_event_corpus,
)
from riskon.event_structure.ir import Document
from riskon.event_structure.matrix import classify_document
from riskon.event_structure.metadata import explicit_metadata, propose_metadata
from riskon.event_structure.parser import parse_html
from riskon.event_structure.records import ConfigurationRecord, RecordProvenance
from riskon.event_structure.references import (
    build_reference_graph,
    load_title_index_xlsx,
)
from riskon.models import Decision, DetectedContext, ManifestEntry, QueryInput, ReasonCode
from riskon.orchestra.source_safety import LocalCorpus
from riskon.provenance import ProvenanceIndex


def _entry(tmp_path: Path, filename: str, title: str, html: str) -> ManifestEntry:
    source = tmp_path / filename
    source.write_text(html, encoding="utf-8")
    return ManifestEntry(
        filename=filename,
        title=title,
        url=f"local://event-wiki/{filename}",
        source_path=str(source),
    )


def _matrix_html() -> str:
    return """<html><body>
    <p>
      <ac:emoticon ac:name="tick"/> Session and Overnight
      <ac:emoticon ac:name="warning"/> only session
    </p>
    <h2>Advisory Location CH (BC CH)</h2>
    <table>
      <tr><th rowspan="2">Service Model</th><th colspan="2">Alerts</th></tr>
      <tr><th>Alert A</th><th>Alert B</th></tr>
      <tr><td>Advice Premium</td><td><ac:emoticon ac:name="tick"/></td>
          <td><ac:image ac:alt="(warning)"><ri:url
          ri:value="https://wiki.juliusbaer.com/icons/emoticons/warning.svg"/>
          </ac:image></td></tr>
      <tr><td>Advice Basic</td><td></td><td><ac:emoticon ac:name="warning"/></td></tr>
    </table>
    <img src="content.png" alt="ordinary content"/>
    </body></html>"""


def test_icons_and_page_legend_keep_raw_state_and_ignore_content_images() -> None:
    document = parse_html(_matrix_html(), "matrix")
    assert document.icon_count == 3
    assert [(entry.icon, entry.meaning) for entry in document.legend.entries] == [
        ("tick", "Session and Overnight"),
        ("warning", "only session"),
    ]
    assert document.tables[0].grid[2][2].icons[0].encoding == "image"
    assert document.tables[0].grid[2][2].icons[0].raw == "(warning)"
    assert not any(issue.code == "unknown_icon" for issue in document.issues)

    image_only = parse_html(
        '<html><body><ac:image ac:alt="(error)"><ri:url '
        'ri:value="/icons/emoticons/error.svg"/></ac:image></body></html>',
        "image",
    )
    assert image_only.tables == []
    assert image_only.icon_count == 0
    assert image_only.issues == []


def test_rowspan_colspan_and_matrix_records_are_governed(tmp_path: Path) -> None:
    document = parse_html(_matrix_html(), "matrix")
    table = document.tables[0]
    assert (table.n_rows, table.n_cols, table.header_rows) == (4, 3, 2)
    assert table.header_labels() == ["Service Model", "Alerts Alert A", "Alerts Alert B"]
    assert table.grid[1][0].spanned
    verdict = classify_document(document)[0]
    assert verdict.is_matrix
    assert verdict.state_columns == [1, 2]

    result = build_structural_event_corpus(
        [_entry(tmp_path, "matrix.html", "Matrix", _matrix_html())]
    )
    matrix_tables = result.tables_for_page("matrix")
    assert len(matrix_tables) == 1
    records = matrix_tables[0].records
    assert len(records) == 4
    premium = next(
        record for record in records if record.dimensions["Service Model"] == "Advice Premium"
    )
    assert premium.state == "tick"
    assert premium.state_meaning == "Session and Overnight"
    assert premium.raw_icon == "tick"
    assert premium.provenance.heading == "Advisory Location CH (BC CH)"


def test_structural_rows_enter_the_existing_provenance_and_retriever(tmp_path: Path) -> None:
    entry = _entry(tmp_path, "matrix.html", "Matrix", _matrix_html())
    sections = ingest_event_sections([entry])
    structural = build_structural_event_corpus([entry])
    sections = attach_structural_tables(sections, structural)
    provenance = ProvenanceIndex(sections, knowledge_root=tmp_path, ref_style="m2")
    structural_units = [unit for unit in provenance.units.values() if unit.structured]
    assert structural_units
    assert all(unit.kind == "table_row" for unit in structural_units)
    assert any("declared meaning:" in unit.text for unit in structural_units)
    config = load_milestone2_config(Path("config/milestone2.toml"))
    retriever = EventHybridRetriever(
        sections,
        provenance,
        config.retrieval,
    )
    assert any(candidate.is_table_row for candidate in retriever.candidates)
    assert any(
        candidate.is_table_row
        and any(provenance.resolve(ref).structured for ref in candidate.unit_refs)
        for candidate in retriever.candidates
    )


def test_reference_graph_reads_xlsx_directly_and_registers_gaps(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["filename", "title", "url"])
    sheet.append(["one.html", "One", "https://wiki.example/one"])
    sheet.append(["two.html", "Two", "https://wiki.example/two"])
    manifest = tmp_path / "manifest.xlsx"
    workbook.save(manifest)

    index = load_title_index_xlsx(manifest)
    assert index == {"one": "one", "two": "two"}
    first = parse_html(
        """<html><body>
        <ri:page ri:content-title="Two"/>
        <ri:page ri:content-title="Missing"/>
        <ac:image><ri:attachment ri:filename="diagram.png"/></ac:image>
        <ac:structured-macro ac:name="children"/>
        </body></html>""",
        "one",
    )
    second = parse_html("<html><body><p>target</p></body></html>", "two")
    graph = build_reference_graph([first, second], title_index=index)
    assert ("one", "two") in graph.edges
    assert any(gap.kind == "unresolved_link" and gap.target == "Missing" for gap in graph.gaps)
    asset_gap = next(gap for gap in graph.gaps if gap.kind == "missing_attachment")
    assert asset_gap.firewall_code == "SOURCE_PACKAGE_ASSET_UNAVAILABLE"
    assert any(gap.kind == "dynamic_content" for gap in graph.gaps)


def test_acronym_registry_verifies_and_surfaces_collisions() -> None:
    assert verify_expansion("OWN", "One-way Notification") == "One-way Notification"
    assert verify_expansion("QQQ", "Questionable Quality") is None
    documents = [
        Document(
            page_id="one",
            source_path="one.html",
            text="Client Investment Profile (CIP). One-way Notification (OWN).",
        ),
        Document(
            page_id="two",
            source_path="two.html",
            text="Customer Investment Product (CIP).",
        ),
    ]
    glossary = build_glossary(documents)
    assert "CIP" in glossary.collisions()
    assert any(entry.acronym == "OWN" and entry.verified for entry in glossary.entries)
    assert not any(entry.acronym == "CIP" for entry in glossary.unique_verified())


def test_declared_scope_can_exclude_but_inferred_scope_cannot() -> None:
    document = Document(
        page_id="scope",
        source_path="scope.html",
        text="Advice Premium applies to BC CH; this is a policy.",
    )
    proposal = propose_metadata(document)
    assert not proposal.metadata.declared
    assert proposal.metadata.covers("BC MC") is None
    declared = explicit_metadata(
        "scope",
        booking_centres=("BC CH",),
        does_not_cover=("Monaco",),
    )
    assert declared.covers("BC CH") is True
    assert declared.covers("Monaco") is False
    assert declared.covers("BC MC") is False


def test_record_schema_preserves_rejected_state_boundaries() -> None:
    provenance = RecordProvenance(page_id="p", table_index=0, row=1, col=1)
    try:
        ConfigurationRecord(
            dimensions={},
            column="Alert",
            state="tick",
            provenance=provenance,
        )
    except ValidationError as error:
        assert "no dimensions" in str(error)
    else:  # pragma: no cover - contract assertion
        raise AssertionError("empty dimensions must be rejected")


def test_missing_source_asset_reaches_the_existing_firewall(tmp_path: Path) -> None:
    html = """<html><body><h1>Diagram page</h1>
    <ac:image><ri:attachment ri:filename="missing.png"/></ac:image>
    </body></html>"""
    entry = _entry(tmp_path, "diagram-page.html", "Diagram page", html)
    structural = build_structural_event_corpus([entry])
    sections = attach_structural_tables(ingest_event_sections([entry]), structural)
    corpus = LocalCorpus(
        sections=tuple(sections),
        provenance=ProvenanceIndex(sections, knowledge_root=tmp_path, ref_style="m2"),
        knowledge_root=tmp_path,
        structural=structural,
    )

    class _Router:
        def route(self, _context: object, _reasons: list[ReasonCode]) -> None:
            return None

    runtime = EventEvidenceReasoningRuntime(  # type: ignore[arg-type]
        corpus,
        object(),
        object(),
        _Router(),
    )
    candidate = SimpleNamespace(
        selected_candidates=(SimpleNamespace(source_ref="local://event-wiki/diagram-page.html"),),
        plan=SimpleNamespace(normalised_query="diagram"),
    )
    decision, _answer, _clarification, reasons, _route, _controls = runtime._firewall(
        QueryInput(query="show the diagram"),
        DetectedContext(),
        ContextAssessment(intent="REFERENCE_LOOKUP"),
        candidate,
        EvidenceAnalysisOutput(
            evidence_sufficiency=EvidenceSufficiencyStatus.VISUAL_REQUIRED,
            material_claims=[],
            unresolved_issues=[],
        ),
        [],
        SupportValidation(),
        [],
        SkepticOutput(objections=[]),
    )
    assert decision is Decision.ABSTAIN
    assert ReasonCode.SOURCE_PACKAGE_ASSET_UNAVAILABLE in reasons
