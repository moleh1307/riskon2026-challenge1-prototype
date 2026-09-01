"""M4D contract freeze: synthetic inputs and policy schemas only."""

import json
import tomllib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
M4D_ROOT = PROJECT_ROOT / "data" / "synthetic" / "m4d"
KNOWLEDGE_ROOT = M4D_ROOT / "knowledge"

EXPECTED_FILES = [
    "ambiguous_arc_registry.html",
    "atlas_critical_control.html",
    "atlas_exception_process.html",
    "meridian_region_alpha.html",
    "meridian_region_beta.html",
    "stability_marker_definition.html",
]
EXPECTED_SIGNALS = [
    "BASELINE_CLARIFY",
    "MISSING_REQUIRED_CONTEXT",
    "AMBIGUOUS_ACRONYM",
    "APPROVAL_REQUIRED",
    "TECHNICAL_FAILURE",
    "UNSUPPORTED_MODALITY",
    "UNRESOLVED_REQUIRED_REFERENCE",
    "CONTRADICTORY_SOURCES",
    "SCOPE_SENSITIVE",
    "JURISDICTION_SENSITIVE",
    "SERVICE_MODEL_SENSITIVE",
    "SOLICITATION_SENSITIVE",
    "WORKFLOW_STAGE_SENSITIVE",
    "COUNTERFACTUAL_REQUIRED",
    "CRITICAL_CONTROL_RISK",
    "TABLE_DEPENDENT",
    "REQUIRED_REFERENCE",
    "MULTI_PART_QUERY",
    "LOW_RETRIEVAL_MARGIN",
    "PROMPT_INJECTION_SIGNAL",
]
EXPECTED_BUDGET = {
    "maximum_total_worker_tasks": 7,
    "maximum_worker_depth": 1,
    "maximum_parallel_discovery_workers": 3,
    "maximum_challenge_rounds": 1,
    "maximum_counterfactual_variants": 3,
    "maximum_counterfactual_depth": 1,
}


def load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    assert isinstance(value, dict)
    return value


def iter_strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [item for child in value.values() for item in iter_strings(child)]
    if isinstance(value, list):
        return [item for child in value for item in iter_strings(child)]
    return []


def test_m4d_file_inventory_and_manifest_are_closed_world() -> None:
    assert sorted(path.name for path in KNOWLEDGE_ROOT.glob("*.html")) == EXPECTED_FILES
    manifest = M4D_ROOT / "manifest.xlsx"
    assert manifest.is_file()
    workbook = load_workbook(manifest, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["manifest"]
        rows = list(workbook["manifest"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows[0] == ("filename", "title", "url")
    assert len(rows) == len(EXPECTED_FILES) + 1
    assert [row[0] for row in rows[1:]] == EXPECTED_FILES
    assert all(row[2] == f"local://synthetic-m4d/{row[0]}" for row in rows[1:])


def test_m4d_policy_schemas_and_bounds_are_frozen() -> None:
    runtime = load_json(M4D_ROOT / "runtime_policy.json")
    assert set(runtime) == {
        "schema_version",
        "activation_policy_ref",
        "counterfactual_policy_ref",
        "auto_activation_enabled",
        "risk_signal_order",
        "risk_signal_rules",
        "execution_budget",
    }
    assert runtime["schema_version"] == "1.0"
    assert runtime["auto_activation_enabled"] is True
    assert runtime["risk_signal_order"] == EXPECTED_SIGNALS
    assert [rule["signal"] for rule in runtime["risk_signal_rules"]] == EXPECTED_SIGNALS
    assert runtime["execution_budget"] == EXPECTED_BUDGET

    failure = load_json(M4D_ROOT / "failure_policy.json")
    assert set(failure) == {"schema_version", "retry_count", "partial_answer_enabled", "rules"}
    assert failure["schema_version"] == "1.0"
    assert failure["retry_count"] == 0
    assert failure["partial_answer_enabled"] is False
    assert {rule["rule_id"] for rule in failure["rules"]} == {
        "ANSWER_WORKER_FAILURE",
        "ABSTAIN_WORKER_FAILURE",
        "CLARIFY_WORKER_FAILURE",
        "CONFIGURATION_FAILURE",
    }


def test_m4d_cases_are_exactly_the_five_declared_activation_cases() -> None:
    document = load_json(M4D_ROOT / "evaluation_cases.json")
    assert set(document) == {"schema_version", "milestone", "cases"}
    assert document["schema_version"] == "1.0"
    cases = document["cases"]
    assert [case["id"] for case in cases] == [
        "M4D-041",
        "M4D-042",
        "M4D-043",
        "M4D-044",
        "M4D-045",
    ]
    assert [case["expected_profile"] for case in cases] == [
        "FAST_PATH",
        "SHORT_CIRCUIT_CLARIFY",
        "HUMAN_FIRST",
        "DUAL_CHECK",
        "FULL_ORCHESTRA",
    ]
    assert cases[1]["expected_final"]["clarifying_question"] == (
        "Do you mean Advisory Review Code or Account Routing Console?"
    )
    assert cases[4]["expected_counterfactuals"][-1]["decision"] == "CLARIFY"


def test_m4d_content_is_synthetic_local_data_and_required_form_is_missing() -> None:
    all_text = " ".join(
        path.read_text(encoding="utf-8") for path in sorted(KNOWLEDGE_ROOT.glob("*.html"))
    )
    assert "http://" not in all_text and "https://" not in all_text
    assert "/Users/" not in all_text and "@" not in all_text
    assert "Bank Julius Baer" not in all_text
    assert "local://synthetic-m4d/attachments/atlas-exception-form.txt" in all_text
    assert not (M4D_ROOT / "attachments" / "atlas-exception-form.txt").exists()
    for token in (
        "Synthetic Stability Marker",
        "Advisory Review Code",
        "Account Routing Console",
        "atlas_do_not_proceed",
        "atlas_client_acceptance_no_override",
        "REGION_ALPHA",
        "REGION_BETA",
        "SERVICE_PLUS",
        "SERVICE_BASIC",
    ):
        assert token in all_text


def test_m4d_config_is_local_only_and_points_to_the_frozen_contract() -> None:
    raw = tomllib.loads((PROJECT_ROOT / "config" / "milestone4d.toml").read_text(encoding="utf-8"))
    assert set(raw) == {"base", "orchestra", "evaluation", "security"}
    assert raw["base"] == {"config": "config/milestone4c.toml"}
    assert raw["orchestra"]["m4d"]["counterfactual_backend"] == "LOCAL_PLANNED_PIPELINE_V1"
    assert raw["orchestra"]["m4d"]["retry_count"] == 0
    assert raw["orchestra"]["m4d"]["partial_answer_enabled"] is False
    assert raw["security"] == {
        "network_enabled": False,
        "external_api_enabled": False,
        "recursive_orchestration_enabled": False,
        "counterfactual_routing_enabled": False,
        "agent_to_agent_citation_enabled": False,
    }
    assert all("http://" not in value and "https://" not in value for value in iter_strings(raw))
