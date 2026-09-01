"""Manifest and HTML preservation tests."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from riskon.ingestion import (
    IngestionError,
    KnowledgeIngestor,
    ManifestLoader,
    load_synthetic_sections,
    manifest_row_count,
)
from riskon.models import ManifestEntry

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = PROJECT_ROOT / "data" / "synthetic"


def test_manifest_has_exact_shape_and_five_rows() -> None:
    assert manifest_row_count(DATA_ROOT / "manifest.xlsx") == (5, 3)
    entries = ManifestLoader().load(DATA_ROOT / "manifest.xlsx", DATA_ROOT)
    assert len(entries) == 5
    assert [entry.url for entry in entries] == [
        "local://synthetic/delegated_order_giver.html",
        "local://synthetic/concentration_workflows.html",
        "local://synthetic/regional_scope_policy.html",
        "local://synthetic/alert_matrix.html",
        "local://synthetic/platform_support.html",
    ]


def test_ingestion_preserves_headings_lists_tables_links_and_image_refs() -> None:
    sections = load_synthetic_sections(DATA_ROOT)
    assert {section.title for section in sections} == {
        "Delegated Order Giver",
        "Concentration Workflows",
        "Regional Scope Policy",
        "Alert Matrix",
        "Platform Support",
    }
    delegated_sections = [
        section for section in sections if section.filename == "delegated_order_giver.html"
    ]
    delegated = delegated_sections[0]
    assert delegated.heading_path == ["Delegated Order Giver"]
    assert any("Product Familiarity Record" in paragraph for paragraph in delegated.paragraphs)
    direct_rule = next(
        section for section in delegated_sections if section.heading_path[-1] == "Direct rule"
    )
    assert direct_rule.lists[0].items == [
        "Check the record.",
        "Keep the instruction within the recorded product scope.",
        "Escalate when the record is missing.",
    ]

    table_section = next(section for section in sections if section.filename == "alert_matrix.html")
    assert table_section.tables[0].headers == ["Alert class", "Status", "Control type"]
    assert table_section.tables[0].rows == [
        ["Concentration", "Open", "Manual review"],
        ["Regional", "Scoped", "Jurisdiction check"],
        ["Technical", "Escalate", "Service desk"],
    ]

    regional = next(
        section for section in sections if section.filename == "regional_scope_policy.html"
    )
    assert regional.links[0].href == "local://synthetic/region-alpha-policy"
    platform = next(section for section in sections if section.filename == "platform_support.html")
    assert platform.images[0].src == "../assets/mock_flow.svg"
    assert platform.images[0].alt == "Synthetic escalation flow diagram"
    assert "Question" not in platform.text


def test_ingestion_flattens_confluence_wrappers_and_custom_images(tmp_path: Path) -> None:
    source = tmp_path / "wrapped.html"
    source.write_text(
        """
        <ac:layout><ac:layout-section><ac:layout-cell>
          <h1>Wrapped page</h1>
          <p>Introductory text.</p>
          <h2>Procedure</h2>
          <ul><li>First step</li><li>Second step</li></ul>
          <table><tr><th>Control</th><th>Status</th></tr>
          <tr><td>Review</td><td>Required</td></tr></table>
          <p><ac:image><ri:attachment ri:filename="flow.png" /></ac:image></p>
        </ac:layout-cell></ac:layout-section></ac:layout>
        """,
        encoding="utf-8",
    )
    sections = KnowledgeIngestor().ingest(
        [
            ManifestEntry(
                filename="wrapped.html",
                title="Wrapped page",
                url="local://event-wiki/wrapped.html",
                source_path=str(source),
            )
        ]
    )

    assert [section.heading_path for section in sections] == [
        ["Wrapped page"],
        ["Wrapped page", "Procedure"],
    ]
    procedure = sections[1]
    assert procedure.lists[0].items == ["First step", "Second step"]
    assert procedure.tables[0].rows == [["Review", "Required"]]
    assert procedure.images[0].src == "attachment/flow.png"


def _write_manifest(path: Path, rows: list[tuple[str, str, str]]) -> None:
    workbook = load_workbook(DATA_ROOT / "manifest.xlsx")
    worksheet = workbook.active
    worksheet.delete_rows(2, worksheet.max_row)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def test_duplicate_filename_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "manifest.xlsx"
    _write_manifest(
        path,
        [
            ("delegated_order_giver.html", "One", "local://synthetic/one.html"),
            ("delegated_order_giver.html", "Two", "local://synthetic/two.html"),
        ],
    )
    with pytest.raises(IngestionError, match="Duplicate"):
        ManifestLoader().load(path, DATA_ROOT)


def test_missing_file_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "manifest.xlsx"
    _write_manifest(path, [("missing.html", "Missing", "local://synthetic/missing.html")])
    with pytest.raises(IngestionError, match="missing"):
        ManifestLoader().load(path, tmp_path)


def test_wrong_manifest_scheme_and_empty_manifest_fail(tmp_path: Path) -> None:
    path = tmp_path / "manifest.xlsx"
    _write_manifest(path, [("delegated_order_giver.html", "Page", "https://not-local.invalid")])
    with pytest.raises(IngestionError, match="outside"):
        ManifestLoader().load(path, DATA_ROOT)
    traversal = tmp_path / "traversal.xlsx"
    _write_manifest(
        traversal,
        [("../delegated_order_giver.html", "Page", "local://synthetic/page.html")],
    )
    with pytest.raises(IngestionError, match="basename"):
        ManifestLoader().load(traversal, DATA_ROOT)
    empty = tmp_path / "empty.xlsx"
    workbook = load_workbook(DATA_ROOT / "manifest.xlsx")
    worksheet = workbook.active
    worksheet.delete_rows(1, worksheet.max_row)
    workbook.save(empty)
    with pytest.raises(IngestionError, match="empty"):
        ManifestLoader().load(empty, DATA_ROOT)
