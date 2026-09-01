"""PowerPoint construction and native-artifact tests for ER-C."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation

from riskon.pitch.catalog import PitchCatalog
from riskon.pitch.deck_builder import _add_flow_row, build_pptx
from riskon.pitch.metrics_loader import load_metric_snapshot


def _slide_text(slide: object) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes  # type: ignore[attr-defined]
        if hasattr(shape, "text") and shape.text
    )


def test_generated_deck_has_twelve_editable_slides_and_notes(event_pitch_build: object) -> None:
    result = event_pitch_build  # type: ignore[assignment]
    presentation = Presentation(str(result.deck_path))  # type: ignore[attr-defined]
    assert len(presentation.slides) == 12
    assert presentation.slide_width / presentation.slide_height == 16 / 9
    for index, slide in enumerate(presentation.slides, start=1):
        text = _slide_text(slide)
        assert text
        assert f"{index:02d}" in text or index == 1
        assert "[Sources]" in slide.notes_slide.notes_text_frame.text


def test_deck_contains_no_media_or_external_relationships(event_pitch_build: object) -> None:
    result = event_pitch_build  # type: ignore[assignment]
    with ZipFile(result.deck_path) as archive:  # type: ignore[attr-defined]
        names = archive.namelist()
        assert not any(name.startswith("ppt/media/") for name in names)
        relationships = b"\n".join(archive.read(name) for name in names if name.endswith(".rels"))
        assert b'TargetMode="External"' not in relationships


def test_build_pptx_is_byte_deterministic(
    event_pitch_catalog: PitchCatalog, event_pitch_config: object, tmp_path: Path
) -> None:
    config = event_pitch_config  # type: ignore[assignment]
    snapshot = load_metric_snapshot(config.runtime.dashboard_metrics)  # type: ignore[attr-defined]
    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"
    build_pptx(event_pitch_catalog, snapshot, first)
    build_pptx(event_pitch_catalog, snapshot, second)
    assert sha256(first.read_bytes()).digest() == sha256(second.read_bytes()).digest()


def test_flow_row_supports_dark_native_shapes() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_flow_row(slide, ["A", "B"], 10, [100, 100], dark=True)
    assert len(slide.shapes) == 5
    assert [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text] == [
        "A",
        "B",
    ]
