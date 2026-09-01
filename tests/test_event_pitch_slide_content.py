"""Content-preservation tests for the generated deck."""

from __future__ import annotations

from pptx import Presentation

from riskon.pitch.catalog import PitchCatalog


def _texts(slide: object) -> list[str]:
    return [
        shape.text
        for shape in slide.shapes  # type: ignore[attr-defined]
        if hasattr(shape, "text") and shape.text
    ]


def test_every_frozen_slide_title_and_key_message_is_in_the_deck(
    event_pitch_build: object, event_pitch_catalog: PitchCatalog
) -> None:
    result = event_pitch_build  # type: ignore[assignment]
    presentation = Presentation(str(result.deck_path))  # type: ignore[attr-defined]
    for slide_number, content in enumerate(event_pitch_catalog.content.slides, start=1):
        slide = presentation.slides[slide_number - 1]
        text = "\n".join(_texts(slide))
        assert content.title in text
        assert content.speaker_key_message in slide.notes_slide.notes_text_frame.text


def test_deck_keeps_main_and_appendix_order(event_pitch_build: object) -> None:
    result = event_pitch_build  # type: ignore[assignment]
    presentation = Presentation(str(result.deck_path))  # type: ignore[attr-defined]
    titles = ["\n".join(_texts(slide)) for slide in presentation.slides]
    assert "RiskON Orchestra" in titles[0]
    assert "APPENDIX" in titles[8]
    assert "APPENDIX" in titles[11]


def test_notes_mark_source_and_key_message(event_pitch_build: object) -> None:
    result = event_pitch_build  # type: ignore[assignment]
    presentation = Presentation(str(result.deck_path))  # type: ignore[attr-defined]
    for slide in presentation.slides:
        notes = slide.notes_slide.notes_text_frame.text
        assert notes.startswith("[Sources]\n- ")
        assert "[Key message]" in notes
