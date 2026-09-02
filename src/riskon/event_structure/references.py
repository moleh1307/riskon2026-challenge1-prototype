"""Manifest-backed internal references and typed source-package gaps."""

from __future__ import annotations

import re
import urllib.parse
from collections import defaultdict
from collections.abc import Sequence
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field

from riskon.event_structure.ir import Document

GapKind = Literal["missing_attachment", "unresolved_link", "dynamic_content"]
WIKI_PATH = re.compile(r"/display/[^/]+/([^?#]+)")


def title_from_wiki_url(url: str) -> str | None:
    match = WIKI_PATH.search(url)
    if not match:
        return None
    title = urllib.parse.unquote(match.group(1).replace("+", " ")).strip()
    return title or None


def load_title_index_xlsx(path: Path) -> dict[str, str]:
    """Read the delivered ``filename,title,url`` workbook directly."""

    workbook = load_workbook(path.expanduser().resolve(), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {}
    headers = [str(value or "").strip().casefold() for value in rows[0]]
    try:
        title_index = headers.index("title")
        filename_index = next(
            headers.index(name) for name in ("filename", "file", "page_id") if name in headers
        )
    except (StopIteration, ValueError) as error:
        raise ValueError("manifest must declare title and filename columns") from error
    result: dict[str, str] = {}
    for row in rows[1:]:
        if title_index >= len(row) or filename_index >= len(row):
            continue
        title = str(row[title_index] or "").strip()
        filename = str(row[filename_index] or "").strip()
        if title and filename:
            result[title.casefold()] = Path(filename).stem
    return result


def title_index_from_pairs(pairs: Sequence[tuple[str, str]]) -> dict[str, str]:
    """Build the same index from already-loaded manifest rows, without conversion."""

    return {title.strip().casefold(): Path(filename).stem for filename, title in pairs if title}


class Gap(BaseModel):
    """A source dependency that cannot be reached from the delivered package."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    kind: GapKind
    target: str
    from_page: str
    detail: str = ""

    @property
    def firewall_code(self) -> str:
        return (
            "SOURCE_PACKAGE_ASSET_UNAVAILABLE" if self.kind == "missing_attachment" else self.kind
        )


class ReferenceGraph(BaseModel):
    """Resolved page edges plus unresolved links, assets, and dynamic macros."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    edges: list[tuple[str, str]] = Field(default_factory=list)
    gaps: list[Gap] = Field(default_factory=list)

    def neighbours(self, page_id: str) -> set[str]:
        return {b for a, b in self.edges if a == page_id} | {
            a for a, b in self.edges if b == page_id
        }

    def gaps_by_kind(self) -> dict[str, int]:
        counts: dict[str, int] = defaultdict(int)
        for gap in self.gaps:
            counts[gap.kind] += 1
        return dict(counts)

    def gaps_for_page(self, page_id: str) -> tuple[Gap, ...]:
        return tuple(gap for gap in self.gaps if gap.from_page == page_id)

    def unresolved_titles(self) -> set[str]:
        return {gap.target for gap in self.gaps if gap.kind == "unresolved_link"}


def build_reference_graph(
    documents: list[Document],
    title_index: dict[str, str] | None = None,
    available_attachments: set[str] | None = None,
) -> ReferenceGraph:
    index = {key.casefold(): value for key, value in (title_index or {}).items()}
    present = available_attachments or set()
    known_pages = {document.page_id for document in documents}
    edges: list[tuple[str, str]] = []
    gaps: list[Gap] = []
    for document in documents:
        for link in document.links:
            if link.kind == "internal_page":
                title = link.target
            elif link.kind == "wiki_url":
                title = title_from_wiki_url(link.target) or ""
            else:
                continue
            if not title:
                continue
            target = index.get(title.casefold())
            if target and target in known_pages:
                edges.append((document.page_id, target))
            else:
                detail = "no page index available" if not index else "title not in page index"
                gaps.append(
                    Gap(
                        kind="unresolved_link",
                        target=title,
                        from_page=document.page_id,
                        detail=detail,
                    )
                )
        for attachment in document.attachments:
            if attachment.filename not in present:
                gaps.append(
                    Gap(
                        kind="missing_attachment",
                        target=attachment.filename,
                        from_page=document.page_id,
                        detail=f"referenced via {attachment.via}, not delivered with the export",
                    )
                )
        for macro in document.macros:
            if macro.dynamic:
                gaps.append(
                    Gap(
                        kind="dynamic_content",
                        target=macro.name,
                        from_page=document.page_id,
                        detail="body generated at render time; not present in source",
                    )
                )
    return ReferenceGraph(edges=sorted(set(edges)), gaps=gaps)
