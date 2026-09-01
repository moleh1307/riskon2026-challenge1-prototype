"""Build the editable ER-C PowerPoint from text, shapes and evaluator data."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from riskon.pitch.catalog import PitchCatalog
from riskon.pitch.metrics_loader import metric_values
from riskon.pitch.models import MetricSnapshot

SLIDE_WIDTH = 1280
SLIDE_HEIGHT = 720
FONT = "Arial"

NAVY = "0B1F33"
WHITE = "FFFFFF"
SAND = "F2E7D5"
GREEN = "2E8B57"
AMBER = "D99000"
RED = "B84545"
INK = "17212B"
MUTED = "5C6B77"
LINE = "D7DEE4"
PALE = "F7F9FA"


def _emu(value: float) -> Any:
    """Convert a design-space pixel value to PowerPoint inches."""

    return Inches(value / 96)


def _rgb(value: str) -> RGBColor:
    """Convert a six-digit hex color to an RGBColor."""

    return cast(RGBColor, RGBColor.from_string(value.replace("#", "")))  # type: ignore[no-untyped-call]


def _add_text(
    slide: Any,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 4,
) -> Any:
    """Add a consistently formatted, editable text box."""

    shape = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(width), _emu(height))
    shape.fill.background()
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = _emu(margin)
    frame.margin_right = _emu(margin)
    frame.margin_top = _emu(margin)
    frame.margin_bottom = _emu(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = _rgb(color)
    return shape


def _add_box(
    slide: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fill: str = WHITE,
    line: str = LINE,
    radius: bool = True,
) -> Any:
    """Add a flat presentation-native panel."""

    geometry = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(geometry, _emu(x), _emu(y), _emu(width), _emu(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line)
    shape.line.width = Pt(0.8)
    return shape


def _add_rule(slide: Any, x: float, y: float, width: float, color: str = LINE) -> None:
    """Add a thin horizontal rule."""

    line = slide.shapes.add_connector(
        1,
        _emu(x),
        _emu(y),
        _emu(x + width),
        _emu(y),
    )
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(0.8)


def _add_chevron(slide: Any, x: float, y: float, width: float, height: float, fill: str) -> Any:
    """Add a small directional chevron between flow nodes."""

    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, _emu(x), _emu(y), _emu(width), _emu(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.fill.background()
    return shape


def _add_circle(slide: Any, x: float, y: float, diameter: float, fill: str) -> Any:
    """Add a filled circle used as a simple visual marker."""

    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, _emu(x), _emu(y), _emu(diameter), _emu(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.fill.background()
    return shape


def _add_header(slide: Any, content: dict[str, Any], number: int, *, dark: bool = False) -> None:
    """Add consistent section label, title and page marker."""

    foreground = WHITE if dark else NAVY
    label_color = SAND if dark else GREEN
    section = "APPENDIX" if number > 8 else "RISKON ORCHESTRA"
    multiline = len(content["title"]) > 50
    _add_text(slide, section, 66, 25, 300, 20, size=11, color=label_color, bold=True, margin=0)
    _add_text(
        slide,
        content["title"],
        66,
        54,
        1140,
        128 if multiline else 82,
        size=35,
        color=foreground,
        bold=True,
        margin=0,
    )
    _add_text(
        slide,
        f"{number:02d}",
        1170,
        28,
        44,
        22,
        size=12,
        color=label_color,
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    _add_rule(slide, 66, 192 if multiline else 139, 1148, SAND if dark else LINE)


def _add_footer(slide: Any, content: dict[str, Any], number: int, *, dark: bool = False) -> None:
    """Add provenance and the synthetic-data disclaimer to every slide."""

    color = "C8D3DB" if dark else MUTED
    _add_text(slide, content["source_label"], 66, 687, 850, 16, size=9, color=color, margin=0)
    _add_text(
        slide,
        "Synthetic local demonstration",
        930,
        687,
        250,
        16,
        size=9,
        color=color,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    _add_text(
        slide, str(number), 1200, 687, 14, 16, size=9, color=color, align=PP_ALIGN.RIGHT, margin=0
    )


def _set_background(slide: Any, color: str) -> None:
    """Set a solid slide background."""

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _add_notes(slide: Any, content: dict[str, Any]) -> None:
    """Add source notes required for every audience-facing slide."""

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "[Sources]\n- "
        + content["source_label"]
        + "\n\n[Key message]\n"
        + content["speaker_key_message"]
    )


def _build_title(slide: Any, content: dict[str, Any]) -> None:
    """Build the deliberately minimal title slide."""

    _set_background(slide, NAVY)
    _add_text(
        slide,
        "RISKON 2026  /  CHALLENGE 1",
        66,
        42,
        440,
        24,
        size=12,
        color=SAND,
        bold=True,
        margin=0,
    )
    _add_text(slide, content["title"], 66, 132, 720, 78, size=58, color=WHITE, bold=True, margin=0)
    _add_text(
        slide,
        "\n".join(content["subtitle_lines"]),
        66,
        235,
        670,
        150,
        size=34,
        color=WHITE,
        bold=True,
        margin=0,
    )
    _add_text(slide, content["supporting_line"], 66, 432, 690, 64, size=22, color=SAND, margin=0)
    _add_rule(slide, 66, 607, 680, SAND)
    _add_text(slide, content["footer"], 66, 624, 650, 22, size=11, color="C8D3DB", margin=0)
    _add_text(
        slide,
        "01",
        1170,
        42,
        44,
        22,
        size=12,
        color=SAND,
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    _add_text(slide, "INVESTIGATE", 922, 170, 230, 26, size=15, color=WHITE, bold=True, margin=0)
    _add_text(
        slide, "independent perspectives", 922, 198, 230, 24, size=14, color="C8D3DB", margin=0
    )
    _add_circle(slide, 858, 171, 38, GREEN)
    _add_text(slide, "EVIDENCE", 922, 315, 230, 26, size=15, color=WHITE, bold=True, margin=0)
    _add_text(slide, "the gate decides", 922, 343, 230, 24, size=14, color="C8D3DB", margin=0)
    _add_circle(slide, 858, 316, 38, SAND)
    _add_text(slide, "APPROVE", 922, 460, 230, 26, size=15, color=WHITE, bold=True, margin=0)
    _add_text(slide, "humans authorise", 922, 488, 230, 24, size=14, color="C8D3DB", margin=0)
    _add_circle(slide, 858, 461, 38, WHITE)
    connector = slide.shapes.add_connector(1, _emu(877), _emu(209), _emu(877), _emu(315))
    connector.line.color.rgb = _rgb("6C8598")
    connector = slide.shapes.add_connector(1, _emu(877), _emu(354), _emu(877), _emu(460))
    connector.line.color.rgb = _rgb("6C8598")


def _build_problem(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the three-stakeholder problem slide."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    for index, card in enumerate(content["cards"]):
        x = 66 + index * 386
        _add_box(slide, x, 218, 360, 210, fill=PALE)
        _add_text(
            slide, f"0{index + 1}", x + 20, 240, 40, 24, size=12, color=GREEN, bold=True, margin=0
        )
        _add_text(
            slide, card["label"], x + 20, 284, 315, 35, size=24, color=NAVY, bold=True, margin=0
        )
        _add_text(slide, card["body"], x + 20, 334, 315, 82, size=18, color=MUTED, margin=0)
    _add_box(slide, 66, 480, 1148, 132, fill=SAND, line=SAND)
    _add_text(
        slide,
        content["bottom_message"],
        94,
        514,
        1090,
        76,
        size=25,
        color=NAVY,
        bold=True,
        margin=0,
    )
    _add_footer(slide, content, number)


def _build_decision(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the four-outcome decision slide."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    positions = [(66, 168), (664, 168), (66, 344), (664, 344)]
    tones = {"green": GREEN, "amber": AMBER, "red": RED, "navy": NAVY}
    for (x, y), outcome in zip(positions, content["outcomes"], strict=True):
        tone = tones[outcome["tone"]]
        _add_box(slide, x, y, 550, 145, fill=WHITE, line=LINE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y), _emu(10), _emu(145))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(tone)
        bar.line.fill.background()
        _add_text(
            slide,
            outcome["label"],
            x + 32,
            y + 24,
            200,
            32,
            size=24,
            color=tone,
            bold=True,
            margin=0,
        )
        _add_text(slide, outcome["body"], x + 32, y + 70, 480, 52, size=18, color=MUTED, margin=0)
    _add_box(slide, 66, 532, 1148, 98, fill=SAND, line=SAND)
    _add_text(
        slide,
        content["bottom_message"],
        92,
        550,
        1100,
        68,
        size=20,
        color=NAVY,
        bold=True,
        margin=0,
    )
    _add_footer(slide, content, number)


def _add_flow_row(
    slide: Any, labels: list[str], y: float, widths: list[float], *, dark: bool = False
) -> None:
    """Draw one bounded flow row with arrows behind the labelled nodes."""

    x: float = 66.0
    fill = NAVY if dark else PALE
    text_color = WHITE if dark else NAVY
    arrow_color = SAND if dark else GREEN
    for index, (label, width) in enumerate(zip(labels, widths, strict=True)):
        if index:
            _add_chevron(slide, x - 18, y + 39, 12, 24, arrow_color)
        _add_box(slide, x, y, width, 104, fill=fill, line=arrow_color)
        _add_text(
            slide,
            label,
            x + 10,
            y + 20,
            width - 20,
            66,
            size=16,
            color=text_color,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        x += width + 24


def _build_architecture(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the two-row Orchestra flow and worker roster."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    _add_flow_row(slide, content["flow"][:4], 216, [150, 250, 110, 170])
    _add_flow_row(slide, content["flow"][4:], 344, [150, 230, 150, 245])
    _add_text(slide, "WORKER ROLES", 66, 482, 220, 22, size=12, color=GREEN, bold=True, margin=0)
    for index, worker in enumerate(content["workers"]):
        x = 66 + index * 231
        _add_box(
            slide,
            x,
            512,
            214,
            72,
            fill=SAND if index == 4 else PALE,
            line=SAND if index == 4 else LINE,
        )
        _add_text(
            slide,
            worker,
            x + 12,
            530,
            190,
            32,
            size=17,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    _add_box(slide, 66, 606, 1148, 54, fill=NAVY, line=NAVY)
    _add_text(
        slide, content["core_rule"], 92, 620, 1090, 26, size=20, color=WHITE, bold=True, margin=0
    )
    _add_footer(slide, content, number)


def _build_demo_interstitial(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the handoff slide for the live demo."""

    _set_background(slide, NAVY)
    _add_text(slide, "RISKON ORCHESTRA", 66, 32, 300, 20, size=11, color=SAND, bold=True, margin=0)
    _add_text(slide, content["title"], 66, 100, 760, 74, size=48, color=WHITE, bold=True, margin=0)
    for index, step in enumerate(content["steps"]):
        y = 230 + index * 112
        _add_circle(slide, 70, y, 42, GREEN if index == 0 else SAND if index == 1 else WHITE)
        _add_text(
            slide,
            str(index + 1),
            70,
            y + 6,
            42,
            28,
            size=18,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            margin=0,
        )
        _add_text(
            slide,
            step,
            136,
            y - 1,
            520,
            55,
            size=20,
            color=WHITE,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    _add_text(
        slide, content["large_line"], 760, 236, 430, 160, size=38, color=SAND, bold=True, margin=0
    )
    _add_rule(slide, 760, 425, 430, SAND)
    _add_text(
        slide,
        "Switch to the frozen local ER-B demo.",
        760,
        446,
        420,
        34,
        size=18,
        color="C8D3DB",
        margin=0,
    )
    _add_text(
        slide,
        "05",
        1170,
        32,
        44,
        22,
        size=12,
        color=SAND,
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    _add_text(slide, content["source_label"], 66, 687, 850, 16, size=9, color="C8D3DB", margin=0)
    _add_text(
        slide,
        "Synthetic local demonstration",
        930,
        687,
        250,
        16,
        size=9,
        color="C8D3DB",
        align=PP_ALIGN.RIGHT,
        margin=0,
    )


def _build_governance(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the governed evolution lifecycle."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    _add_flow_row(slide, content["lifecycle"], 164, [145, 150, 170, 110, 170, 185, 140])
    for index, guardrail in enumerate(content["guardrails"]):
        x = 66 + index * 231
        fill = SAND if index == 3 or index == 4 else PALE
        line = AMBER if index == 3 else RED if index == 4 else LINE
        _add_box(slide, x, 340, 214, 126, fill=fill, line=line)
        _add_text(
            slide,
            guardrail,
            x + 14,
            369,
            186,
            68,
            size=19,
            color=RED if index >= 3 else NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    _add_box(slide, 66, 520, 1148, 112, fill=NAVY, line=NAVY)
    _add_text(
        slide, content["core_message"], 94, 542, 1090, 74, size=21, color=WHITE, bold=True, margin=0
    )
    _add_footer(slide, content, number)


def _build_metrics(
    slide: Any, content: dict[str, Any], snapshot: MetricSnapshot, number: int
) -> None:
    """Build evaluator-backed metric tiles with values read at build time."""

    _set_background(slide, NAVY)
    _add_header(slide, content, number, dark=True)
    ids = content["metric_ids"]
    values = metric_values(snapshot, ids)
    by_id = {metric.id: metric for metric in snapshot.metrics}
    positions = [(66, 230), (354, 230), (642, 230), (930, 230), (66, 410), (354, 410), (642, 410)]
    for metric_id, (x, y) in zip(ids, positions, strict=True):
        metric = by_id[metric_id]
        tone = GREEN if metric_id not in {"unsafe_routing"} else RED
        _add_box(slide, x, y, 258, 154, fill=WHITE, line=WHITE)
        _add_text(
            slide,
            values[metric_id],
            x + 18,
            y + 24,
            220,
            52,
            size=34,
            color=tone,
            bold=True,
            margin=0,
        )
        _add_text(
            slide, metric.label, x + 18, y + 92, 220, 42, size=17, color=NAVY, bold=True, margin=0
        )
    _add_text(slide, content["metric_footer"], 66, 590, 600, 26, size=16, color=SAND, margin=0)
    _add_text(
        slide,
        "Values are sourced from the current generated dashboard metrics.",
        66,
        622,
        790,
        24,
        size=15,
        color="C8D3DB",
        margin=0,
    )
    _add_footer(slide, content, number, dark=True)


def _build_value(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the business-value close."""

    _set_background(slide, SAND)
    _add_header(slide, content, number)
    for index, pillar in enumerate(content["pillars"]):
        x = 66 + index * 386
        _add_box(slide, x, 220, 360, 210, fill=WHITE, line=WHITE)
        _add_text(
            slide, f"0{index + 1}", x + 20, 242, 40, 24, size=12, color=GREEN, bold=True, margin=0
        )
        _add_text(
            slide, pillar["label"], x + 20, 284, 310, 38, size=22, color=NAVY, bold=True, margin=0
        )
        _add_text(slide, pillar["body"], x + 20, 340, 310, 82, size=17, color=MUTED, margin=0)
    _add_box(slide, 66, 450, 1148, 174, fill=NAVY, line=NAVY)
    _add_text(
        slide,
        content["final_statement"],
        96,
        480,
        1085,
        116,
        size=24,
        color=WHITE,
        bold=True,
        margin=0,
    )
    _add_footer(slide, content, number)


def _build_routing(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build role-first functional routing appendix."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    _add_flow_row(slide, content["route_flow"], 166, [188, 188, 188, 188, 260])
    _add_text(
        slide,
        "PERSON-LEVEL SELECTION FACTORS",
        66,
        350,
        430,
        22,
        size=12,
        color=GREEN,
        bold=True,
        margin=0,
    )
    _add_text(
        slide,
        "  ·  ".join(content["factors"]),
        66,
        389,
        1120,
        50,
        size=23,
        color=NAVY,
        bold=True,
        margin=0,
    )
    _add_box(slide, 66, 490, 1148, 94, fill=SAND, line=SAND)
    _add_text(
        slide,
        content["hard_constraint_note"],
        94,
        520,
        1084,
        38,
        size=22,
        color=NAVY,
        bold=True,
        margin=0,
    )
    _add_footer(slide, content, number)


def _build_security(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the local-first security appendix."""

    _set_background(slide, NAVY)
    _add_header(slide, content, number, dark=True)
    for index, control in enumerate(content["controls"]):
        col = index % 2
        row = index // 2
        x = 92 + col * 558
        y = 174 + row * 76
        _add_circle(slide, x, y + 3, 28, GREEN if "No " not in control else SAND)
        _add_text(
            slide,
            "✓",
            x,
            y + 5,
            28,
            22,
            size=15,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            margin=0,
        )
        _add_text(
            slide,
            control,
            x + 48,
            y - 2,
            450,
            38,
            size=22,
            color=WHITE,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    _add_box(slide, 66, 548, 1148, 76, fill=SAND, line=SAND)
    _add_text(slide, content["footer"], 92, 570, 1090, 36, size=19, color=NAVY, bold=True, margin=0)
    _add_footer(slide, content, number, dark=True)


def _build_event_adaptation(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the ER-A event-corpus adaptation appendix."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    _add_flow_row(slide, content["intake_flow"], 220, [185, 185, 220, 220, 200])
    _add_text(slide, "ADAPTER STATUS", 66, 382, 260, 22, size=12, color=GREEN, bold=True, margin=0)
    status_colors = [GREEN, AMBER, RED]
    for index, status in enumerate(content["statuses"]):
        x = 66 + index * 386
        _add_box(slide, x, 420, 360, 86, fill=PALE, line=status_colors[index])
        _add_text(
            slide,
            status,
            x + 16,
            444,
            328,
            36,
            size=18,
            color=status_colors[index],
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    _add_box(slide, 66, 538, 1148, 100, fill=SAND, line=SAND)
    _add_text(
        slide, content["message"], 92, 556, 1090, 66, size=19, color=NAVY, bold=True, margin=0
    )
    _add_footer(slide, content, number)


def _build_limitations(slide: Any, content: dict[str, Any], number: int) -> None:
    """Build the honest capability-boundary appendix."""

    _set_background(slide, WHITE)
    _add_header(slide, content, number)
    _add_box(slide, 66, 164, 550, 310, fill=PALE, line=GREEN)
    _add_box(slide, 664, 164, 550, 310, fill=PALE, line=RED)
    _add_text(slide, "PROVES", 94, 190, 200, 30, size=24, color=GREEN, bold=True, margin=0)
    _add_text(
        slide,
        "\n".join("• " + value for value in content["proves"]),
        94,
        238,
        480,
        210,
        size=17,
        color=INK,
        margin=0,
    )
    _add_text(
        slide, "DOES NOT YET PROVE", 692, 190, 480, 30, size=24, color=RED, bold=True, margin=0
    )
    _add_text(
        slide,
        "\n".join("• " + value for value in content["does_not_prove"]),
        692,
        238,
        480,
        210,
        size=17,
        color=INK,
        margin=0,
    )
    _add_box(slide, 66, 520, 1148, 114, fill=NAVY, line=NAVY)
    _add_text(
        slide,
        "\n".join(content["critical_wording"]),
        94,
        538,
        1090,
        84,
        size=17,
        color=WHITE,
        bold=True,
        margin=0,
    )
    _add_footer(slide, content, number)


def build_pptx(catalog: PitchCatalog, snapshot: MetricSnapshot, output: Path) -> None:
    """Build and save the twelve-slide editable deck."""

    prs = Presentation()
    prs.slide_width = _emu(SLIDE_WIDTH)
    prs.slide_height = _emu(SLIDE_HEIGHT)
    blank = prs.slide_layouts[6]
    builders = {
        1: lambda slide, content: _build_title(slide, content),
        2: lambda slide, content: _build_problem(slide, content, 2),
        3: lambda slide, content: _build_decision(slide, content, 3),
        4: lambda slide, content: _build_architecture(slide, content, 4),
        5: lambda slide, content: _build_demo_interstitial(slide, content, 5),
        6: lambda slide, content: _build_governance(slide, content, 6),
        7: lambda slide, content: _build_metrics(slide, content, snapshot, 7),
        8: lambda slide, content: _build_value(slide, content, 8),
        9: lambda slide, content: _build_routing(slide, content, 9),
        10: lambda slide, content: _build_security(slide, content, 10),
        11: lambda slide, content: _build_event_adaptation(slide, content, 11),
        12: lambda slide, content: _build_limitations(slide, content, 12),
    }
    for number in range(1, 13):
        slide = prs.slides.add_slide(blank)
        content = catalog.slide(number).model_dump(mode="json")
        builders[number](slide, content)
        _add_notes(slide, content)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
