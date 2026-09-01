"""Closed-world ingestion for the synthetic XLSX manifest and HTML corpus."""

from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook  # type: ignore[import-untyped]

from riskon.models import (
    ImageRef,
    LinkRef,
    ListBlock,
    ManifestEntry,
    Section,
    TableData,
)


class IngestionError(ValueError):
    """Raised when the closed-world corpus cannot be trusted."""


class ManifestLoader:
    """Load and validate the exact three-column local manifest."""

    expected_headers = ("filename", "title", "url")

    def load(
        self,
        manifest_path: Path,
        data_root: Path,
        *,
        knowledge_root: Path | None = None,
        url_prefix: str = "local://synthetic/",
    ) -> list[ManifestEntry]:
        if not manifest_path.is_file():
            raise IngestionError(f"Manifest not found: {manifest_path}")

        workbook = load_workbook(manifest_path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not rows:
            raise IngestionError("Manifest is empty")
        headers = tuple(str(value).strip() if value is not None else "" for value in rows[0])
        if headers != self.expected_headers:
            raise IngestionError(f"Manifest headers must be {self.expected_headers}")

        entries: list[ManifestEntry] = []
        seen: set[str] = set()
        for row_number, row in enumerate(rows[1:], start=2):
            values = list(row[:3])
            if not values or all(value is None for value in values):
                continue
            if len(values) != 3 or any(value is None for value in values):
                raise IngestionError(f"Manifest row {row_number} is incomplete")
            filename, title, url = (str(value).strip() for value in values)
            if not filename or not title or not url:
                raise IngestionError(f"Manifest row {row_number} has an empty value")
            if filename in seen:
                raise IngestionError(f"Duplicate manifest filename: {filename}")
            if not filename.endswith(".html"):
                raise IngestionError(f"Only local HTML is allowed in M0: {filename}")
            filename_path = Path(filename)
            if filename_path.is_absolute() or filename_path.name != filename:
                raise IngestionError(f"Manifest filename must be a local basename: {filename}")
            if not url.startswith(url_prefix):
                raise IngestionError(f"Manifest URL is outside the synthetic scheme: {url}")

            resolved_knowledge_root = (knowledge_root or (data_root / "knowledge")).resolve()
            source_path = (resolved_knowledge_root / filename).resolve()
            if resolved_knowledge_root not in source_path.parents:
                raise IngestionError(f"Manifest file is outside the synthetic corpus: {filename}")
            if not source_path.is_file():
                raise IngestionError(f"Manifest file is missing: {source_path}")
            seen.add(filename)
            entries.append(
                ManifestEntry(
                    filename=filename,
                    title=title,
                    url=url,
                    source_path=str(source_path),
                )
            )

        if not entries:
            raise IngestionError("Manifest contains no data rows")
        return entries


@dataclass
class _SectionBuilder:
    """Mutable internal accumulator used while walking one HTML document."""

    heading_path: list[str]
    paragraphs: list[str] = field(default_factory=list)
    lists: list[ListBlock] = field(default_factory=list)
    tables: list[TableData] = field(default_factory=list)
    links: list[LinkRef] = field(default_factory=list)
    images: list[ImageRef] = field(default_factory=list)
    fragments: list[str] = field(default_factory=list)
    scope: dict[str, str] = field(default_factory=dict)
    intents: list[str] = field(default_factory=list)
    claims: dict[str, str] = field(default_factory=dict)


class KnowledgeIngestor:
    """Parse HTML while preserving the structures needed by later milestones."""

    def ingest(self, entries: list[ManifestEntry]) -> list[Section]:
        sections: list[Section] = []
        for entry in entries:
            sections.extend(self._ingest_entry(entry))
        return sections

    def _ingest_entry(self, entry: ManifestEntry) -> list[Section]:
        source_path = Path(entry.source_path)
        raw_html = source_path.read_text(encoding="utf-8")
        soup = BeautifulSoup(raw_html, "lxml")
        document_title = soup.title.get_text(" ", strip=True) if soup.title else entry.title
        body = soup.body or soup
        builders: list[_SectionBuilder] = []
        current: _SectionBuilder | None = None
        heading_stack: list[str] = []

        for child in body.find_all(recursive=False):
            if not isinstance(child, Tag):
                continue
            name = child.name.lower()
            if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                level = int(name[1])
                heading = self._text(child)
                heading_stack = heading_stack[: level - 1] + [heading]
                current = _SectionBuilder(heading_path=list(heading_stack))
                builders.append(current)
                current.fragments.append(heading)
                self._capture_metadata(child, current)
                continue

            if current is None:
                current = _SectionBuilder(heading_path=[])
                builders.append(current)
            self._consume_element(child, current)

        if not builders:
            builders.append(_SectionBuilder(heading_path=[], fragments=[document_title]))

        links = self._unique_links(soup)
        images = self._unique_images(soup)
        for builder in builders:
            builder.links = links
            builder.images = images

        result: list[Section] = []
        for index, builder in enumerate(builders, start=1):
            result.append(
                Section(
                    section_id=f"{entry.filename}::section-{index:02d}",
                    filename=entry.filename,
                    title=document_title,
                    source_ref=entry.url,
                    heading_path=builder.heading_path,
                    text="\n".join(fragment for fragment in builder.fragments if fragment),
                    paragraphs=builder.paragraphs,
                    lists=builder.lists,
                    tables=builder.tables,
                    links=builder.links,
                    images=builder.images,
                    scope=builder.scope,
                    intents=builder.intents,
                    claims=builder.claims,
                )
            )
        return result

    def _consume_element(self, element: Tag, builder: _SectionBuilder) -> None:
        self._capture_metadata(element, builder)
        name = element.name.lower()
        if name == "p":
            text = self._text(element)
            if text:
                builder.paragraphs.append(text)
                builder.fragments.append(text)
                claim_id = self._attribute_text(element, "data-claim-id")
                if claim_id:
                    builder.claims[claim_id] = text
            return
        if name in {"ul", "ol"}:
            items = [self._text(item) for item in element.find_all("li", recursive=False)]
            items = [item for item in items if item]
            if items:
                builder.lists.append(ListBlock(kind=name, items=items))
                builder.fragments.append("\n".join(f"- {item}" for item in items))
            return
        if name == "table":
            table = self._table(element)
            builder.tables.append(table)
            for table_claim_id, row in zip(table.claim_ids, table.rows, strict=False):
                if table_claim_id:
                    builder.claims[table_claim_id] = "; ".join(
                        f"{header}: {value}"
                        for header, value in zip(table.headers, row, strict=False)
                    )
            builder.fragments.append(
                "\n".join([" | ".join(table.headers)] + [" | ".join(row) for row in table.rows])
            )
            return
        text = self._text(element)
        if text:
            builder.fragments.append(text)
        for image in element.find_all("img"):
            alt = str(image.get("alt", "")).strip()
            if alt:
                builder.fragments.append(f"Image reference: {alt}")

    def _table(self, table: Tag) -> TableData:
        rows: list[list[str]] = []
        headers: list[str] = []
        claim_ids: list[str | None] = []
        row_scopes: list[dict[str, str]] = []
        header_row = table.find("thead")
        if header_row:
            header_cells = header_row.find_all(["th", "td"])
            headers = [self._text(cell) for cell in header_cells]
        for row in table.find_all("tr"):
            cells = [self._text(cell) for cell in row.find_all(["th", "td"], recursive=False)]
            if not cells:
                continue
            if not headers and row.find("th"):
                headers = cells
                continue
            if cells != headers:
                rows.append(cells)
                claim_ids.append(self._attribute_text(row, "data-claim-id") or None)
                row_scopes.append(self._scope_attrs(row))
        return TableData(
            headers=headers,
            rows=rows,
            claim_ids=claim_ids,
            row_scopes=row_scopes,
        )

    @classmethod
    def _capture_metadata(cls, element: Tag, builder: _SectionBuilder) -> None:
        builder.scope.update(cls._scope_attrs(element))
        intent = cls._attribute_text(element, "data-intent")
        if intent and intent not in builder.intents:
            builder.intents.append(intent)

    @staticmethod
    def _attribute_text(element: Tag, name: str) -> str:
        value = element.get(name)
        if isinstance(value, list):
            return " ".join(str(item) for item in value).strip()
        return str(value).strip() if value is not None else ""

    @classmethod
    def _scope_attrs(cls, element: Tag) -> dict[str, str]:
        scope: dict[str, str] = {}
        for name, _value in element.attrs.items():
            if not name.startswith("data-scope-"):
                continue
            key = name.removeprefix("data-scope-").replace("-", "_")
            scope[key] = cls._attribute_text(element, name).upper().replace(" ", "_")
        return scope

    @staticmethod
    def _text(tag: Tag) -> str:
        return " ".join(tag.get_text(" ", strip=True).split())

    @staticmethod
    def _unique_links(soup: BeautifulSoup) -> list[LinkRef]:
        links: list[LinkRef] = []
        seen: set[tuple[str, str]] = set()
        for tag in soup.find_all("a"):
            href = str(tag.get("href", "")).strip()
            text = " ".join(tag.get_text(" ", strip=True).split())
            key = (text, href)
            if href and key not in seen:
                links.append(LinkRef(text=text, href=href))
                seen.add(key)
        return links

    @staticmethod
    def _unique_images(soup: BeautifulSoup) -> list[ImageRef]:
        images: list[ImageRef] = []
        seen: set[tuple[str, str]] = set()
        for tag in soup.find_all("img"):
            src = str(tag.get("src", "")).strip()
            alt = str(tag.get("alt", "")).strip()
            key = (src, alt)
            if src and key not in seen:
                images.append(ImageRef(src=src, alt=alt))
                seen.add(key)
        return images


def load_synthetic_sections(data_root: Path) -> list[Section]:
    """Convenience loader used by the pipeline and tests."""

    manifest = ManifestLoader().load(data_root / "manifest.xlsx", data_root)
    return KnowledgeIngestor().ingest(manifest)


def load_sections_from_manifest(
    manifest_path: Path,
    knowledge_root: Path,
    *,
    url_prefix: str,
) -> list[Section]:
    """Load a local corpus whose manifest and knowledge directory are separate."""

    entries = ManifestLoader().load(
        manifest_path,
        knowledge_root.parent,
        knowledge_root=knowledge_root,
        url_prefix=url_prefix,
    )
    return KnowledgeIngestor().ingest(entries)


def manifest_row_count(manifest_path: Path) -> tuple[int, int]:
    """Return data-row and column counts for a lightweight acceptance check."""

    workbook = load_workbook(manifest_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    return max(0, len(rows) - 1), len(rows[0]) if rows else 0
