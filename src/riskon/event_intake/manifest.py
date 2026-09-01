"""Safe Excel manifest loading for event-day HTML corpora."""

from __future__ import annotations

import json
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from urllib.parse import quote

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from riskon.event_intake.errors import ManifestError
from riskon.event_intake.models import (
    CorpusIssue,
    CorpusSmokeCase,
    IssueSeverity,
    ManifestDescriptor,
    ManifestRow,
)

DEFAULT_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "filename": ("filename", "file_name", "file"),
    "title": ("title", "page_title", "page title"),
    "url": ("url", "page_url", "page url"),
}


@dataclass(frozen=True)
class EventManifestEntry:
    """One valid manifest row plus its resolved source path."""

    row: ManifestRow
    source_path: Path
    source_url: str


@dataclass(frozen=True)
class ManifestLoadResult:
    """Manifest descriptor, usable rows, and safe issues."""

    manifest: ManifestDescriptor | None
    entries: tuple[EventManifestEntry, ...]
    issues: tuple[CorpusIssue, ...]


class ManifestLoader:
    """Load an XLSX workbook without evaluating macros, formulas, or links."""

    def __init__(self, aliases_path: Path | None = None) -> None:
        self.aliases = load_column_aliases(aliases_path)

    def load(
        self,
        manifest_path: Path,
        source_root: Path,
        *,
        column_mapping: Mapping[str, str] | None = None,
    ) -> ManifestLoadResult:
        """Read and validate a workbook using only the closed alias registry."""

        root = source_root.resolve(strict=False)
        manifest = manifest_path.resolve(strict=False)
        relative_manifest = (
            manifest.relative_to(root).as_posix()
            if manifest.is_relative_to(root)
            else "<external-manifest>"
        )
        if not manifest.is_file():
            return ManifestLoadResult(
                manifest=None,
                entries=(),
                issues=(
                    CorpusIssue(
                        code="MANIFEST_NOT_FOUND",
                        severity=IssueSeverity.ERROR,
                        relative_path=relative_manifest,
                        detail="The manifest is unavailable at the declared path.",
                    ),
                ),
            )
        if manifest.suffix.lower() != ".xlsx":
            return ManifestLoadResult(
                manifest=None,
                entries=(),
                issues=(
                    CorpusIssue(
                        code="MANIFEST_SCHEMA_UNRECOGNISED",
                        severity=IssueSeverity.ERROR,
                        relative_path=relative_manifest,
                        detail="Only non-macro XLSX manifests are supported.",
                    ),
                ),
            )

        issues: list[CorpusIssue] = []
        if _has_external_workbook_links(manifest):
            issues.append(
                CorpusIssue(
                    code="EXTERNAL_LINK_PRESENT",
                    severity=IssueSeverity.WARNING,
                    relative_path=relative_manifest,
                    detail="External workbook links were observed and not opened.",
                )
            )

        try:
            workbook = load_workbook(
                manifest,
                read_only=True,
                data_only=True,
                keep_links=False,
                keep_vba=False,
            )
        except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
            del exc
            issues.append(
                CorpusIssue(
                    code="MANIFEST_SCHEMA_UNRECOGNISED",
                    severity=IssueSeverity.ERROR,
                    relative_path=relative_manifest,
                    detail="The workbook could not be read as a safe XLSX manifest.",
                )
            )
            return ManifestLoadResult(None, (), tuple(issues))

        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_values = next(rows)
            except StopIteration:
                header_values = ()
            headers = [_cell_text(value) for value in header_values]
            mapping, mapping_issues = self._resolve_mapping(headers, column_mapping)
            issues.extend(mapping_issues)
            if mapping is None:
                return ManifestLoadResult(
                    None,
                    (),
                    tuple(_unique_issues(issues)),
                )

            header_indexes = {header: index for index, header in enumerate(headers)}
            manifest_rows: list[ManifestRow] = []
            entries: list[EventManifestEntry] = []
            seen_filenames: dict[str, int] = {}
            seen_titles: dict[str, int] = {}
            row_count = 0
            for row_number, row in enumerate(rows, start=2):
                values = list(row)
                if not values or all(value is None or not str(value).strip() for value in values):
                    continue
                row_count += 1
                filename = _cell_text(_value_at(row, header_indexes[mapping["filename"]]))
                title = _cell_text(_value_at(row, header_indexes[mapping["title"]]))
                url = (
                    _cell_text(_value_at(row, header_indexes[mapping["url"]]))
                    if mapping.get("url") is not None
                    else ""
                )
                if not filename or not title:
                    issues.append(
                        CorpusIssue(
                            code="MANIFEST_ROW_INVALID",
                            severity=IssueSeverity.ERROR,
                            relative_path=relative_manifest,
                            detail=f"Manifest row {row_number} lacks a filename or title.",
                        )
                    )
                    continue
                safe_filename = _normalize_filename(filename)
                if safe_filename is None:
                    issues.append(
                        CorpusIssue(
                            code="MANIFEST_SCHEMA_UNRECOGNISED",
                            severity=IssueSeverity.ERROR,
                            relative_path=relative_manifest,
                            detail=f"Manifest row {row_number} has an unsafe filename.",
                        )
                    )
                    continue
                row_model = ManifestRow(
                    row_number=row_number,
                    filename=safe_filename,
                    title=title,
                    url=url or None,
                )
                manifest_rows.append(row_model)
                filename_key = safe_filename.casefold()
                duplicate_filename = filename_key in seen_filenames
                if duplicate_filename:
                    issues.append(
                        CorpusIssue(
                            code="DUPLICATE_FILENAME",
                            severity=IssueSeverity.ERROR,
                            relative_path=safe_filename,
                            detail="The manifest declares the same filename more than once.",
                        )
                    )
                else:
                    seen_filenames[filename_key] = row_number
                title_key = title.casefold()
                if title_key in seen_titles:
                    issues.append(
                        CorpusIssue(
                            code="DUPLICATE_TITLE",
                            severity=IssueSeverity.WARNING,
                            relative_path=relative_manifest,
                            detail="The manifest declares a duplicate page title.",
                        )
                    )
                else:
                    seen_titles[title_key] = row_number

                source_path = (root / safe_filename).resolve(strict=False)
                if not source_path.is_relative_to(root) or not source_path.is_file():
                    issues.append(
                        CorpusIssue(
                            code="MANIFEST_HTML_MISSING",
                            severity=IssueSeverity.ERROR,
                            relative_path=safe_filename,
                            detail="A manifest row does not resolve to a local HTML file.",
                        )
                    )
                    continue
                if duplicate_filename:
                    continue
                source_url = _local_source_url(safe_filename, url)
                entries.append(
                    EventManifestEntry(
                        row=row_model,
                        source_path=source_path,
                        source_url=source_url,
                    )
                )
                if url and _is_external_url(url):
                    issues.append(
                        CorpusIssue(
                            code="EXTERNAL_LINK_PRESENT",
                            severity=IssueSeverity.WARNING,
                            relative_path=relative_manifest,
                            detail="An external manifest URL was recorded without fetching it.",
                        )
                    )

            if not manifest_rows:
                issues.append(
                    CorpusIssue(
                        code="MANIFEST_SCHEMA_UNRECOGNISED",
                        severity=IssueSeverity.ERROR,
                        relative_path=relative_manifest,
                        detail="The manifest contains no usable data rows.",
                    )
                )
            descriptor = ManifestDescriptor(
                relative_path=relative_manifest,
                columns=headers,
                column_mapping=mapping,
                row_count=row_count,
                has_url=mapping.get("url") is not None,
                rows=manifest_rows,
            )
            return ManifestLoadResult(
                descriptor,
                tuple(entries),
                tuple(_unique_issues(issues)),
            )
        finally:
            workbook.close()

    def _resolve_mapping(
        self,
        headers: list[str],
        explicit: Mapping[str, str] | None,
    ) -> tuple[dict[str, str] | None, list[CorpusIssue]]:
        """Resolve logical columns using explicit mapping or frozen aliases."""

        issues: list[CorpusIssue] = []
        header_by_key: dict[str, str] = {}
        for header in headers:
            key = _header_key(header)
            if not key:
                continue
            if key in header_by_key:
                issues.append(
                    CorpusIssue(
                        code="MANIFEST_SCHEMA_UNRECOGNISED",
                        severity=IssueSeverity.ERROR,
                        detail="The workbook contains duplicate normalized headers.",
                    )
                )
            header_by_key[key] = header

        supplied = dict(explicit or {})
        unknown = sorted(set(supplied) - set(self.aliases))
        if unknown:
            issues.append(
                CorpusIssue(
                    code="MANIFEST_SCHEMA_UNRECOGNISED",
                    severity=IssueSeverity.ERROR,
                    detail="Explicit mapping contains an unknown logical column.",
                )
            )
            return None, issues

        resolved: dict[str, str] = {}
        for logical in ("filename", "title", "url"):
            requested = supplied.get(logical)
            if requested is not None:
                actual = header_by_key.get(_header_key(requested))
                if actual is None:
                    issues.append(
                        CorpusIssue(
                            code="MANIFEST_SCHEMA_UNRECOGNISED",
                            severity=IssueSeverity.ERROR,
                            detail=f"Explicit mapping for {logical} is not present.",
                        )
                    )
                    continue
                resolved[logical] = actual
                continue
            alias_keys = {_header_key(alias) for alias in self.aliases[logical]}
            matches = [header for key, header in header_by_key.items() if key in alias_keys]
            if len(matches) == 1:
                resolved[logical] = matches[0]
            elif len(matches) > 1:
                issues.append(
                    CorpusIssue(
                        code="MANIFEST_SCHEMA_UNRECOGNISED",
                        severity=IssueSeverity.ERROR,
                        detail=f"Multiple headers match logical column {logical}.",
                    )
                )
            elif logical in {"filename", "title"}:
                issues.append(
                    CorpusIssue(
                        code="MANIFEST_SCHEMA_UNRECOGNISED",
                        severity=IssueSeverity.ERROR,
                        detail=f"Required logical column {logical} was not recognized.",
                    )
                )
        if any(item.severity is IssueSeverity.ERROR for item in issues):
            return None, issues
        return resolved, issues


def load_column_aliases(path: Path | None) -> dict[str, tuple[str, ...]]:
    """Load and validate the small versioned alias registry."""

    if path is None:
        return dict(DEFAULT_COLUMN_ALIASES)
    resolved = path.resolve(strict=False)
    try:
        payload = json.loads(resolved.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ManifestError("Manifest alias registry could not be read") from exc
    if not isinstance(payload, dict) or set(payload) != {"schema_version", "columns"}:
        raise ManifestError("Manifest alias registry schema is not recognized")
    if payload.get("schema_version") != "1.0" or not isinstance(payload.get("columns"), dict):
        raise ManifestError("Manifest alias registry schema is not recognized")
    columns = payload["columns"]
    if set(columns) != set(DEFAULT_COLUMN_ALIASES):
        raise ManifestError("Manifest alias registry columns are not recognized")
    result: dict[str, tuple[str, ...]] = {}
    for logical, values in columns.items():
        if (
            not isinstance(values, list)
            or not values
            or not all(isinstance(value, str) and value.strip() for value in values)
        ):
            raise ManifestError("Manifest alias registry contains invalid aliases")
        result[str(logical)] = tuple(str(value).strip() for value in values)
    return result


def load_smoke_cases(path: Path) -> tuple[CorpusSmokeCase, ...]:
    """Load a closed smoke-case JSON envelope."""

    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ManifestError("Smoke-case file could not be read") from exc
    if not isinstance(payload, dict) or set(payload) != {"schema_version", "cases"}:
        raise ManifestError("Smoke-case schema is not recognized")
    if payload.get("schema_version") != "1.0" or not isinstance(payload.get("cases"), list):
        raise ManifestError("Smoke-case schema is not recognized")
    try:
        cases = tuple(CorpusSmokeCase.model_validate(item) for item in payload["cases"])
    except (TypeError, ValueError) as exc:
        raise ManifestError("Smoke-case schema is not recognized") from exc
    if not cases:
        raise ManifestError("Smoke-case file contains no cases")
    return cases


def _value_at(row: tuple[object, ...], index: int) -> object:
    """Return a cell by its zero-based worksheet index."""

    return row[index] if index < len(row) else None


def _cell_text(value: object) -> str:
    """Normalize a workbook cell without interpreting formulas."""

    return str(value).strip() if value is not None else ""


def _header_key(value: str) -> str:
    """Normalize case, spacing, and underscore aliases."""

    return " ".join(value.replace("_", " ").split()).casefold()


def _normalize_filename(value: str) -> str | None:
    """Accept a relative HTML path while rejecting traversal and drive paths."""

    normalized = value.replace("\\", "/").strip()
    if not normalized or normalized.startswith("/") or ":" in normalized.split("/", 1)[0]:
        return None
    path = PurePosixPath(normalized)
    if any(part in {"", ".", ".."} for part in path.parts):
        return None
    if path.suffix.casefold() not in {".html", ".htm"}:
        return None
    return path.as_posix()


def _local_source_url(relative: str, manifest_url: str) -> str:
    """Keep provenance local even when the workbook contains an external URL."""

    if manifest_url.startswith("local://"):
        return manifest_url
    return f"local://event-corpus/{quote(relative, safe='/')}"


def _is_external_url(value: str) -> bool:
    """Return whether a value would require external fetching."""

    return value.casefold().startswith(("http://", "https://", "//"))


def _has_external_workbook_links(path: Path) -> bool:
    """Detect workbook external-link parts without opening or resolving them."""

    try:
        with zipfile.ZipFile(path) as archive:
            return any(name.startswith("xl/externalLinks/") for name in archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False


def _unique_issues(issues: list[CorpusIssue]) -> list[CorpusIssue]:
    """Remove duplicate issue rows while preserving deterministic order."""

    seen: set[tuple[str, IssueSeverity, str | None, str | None]] = set()
    unique: list[CorpusIssue] = []
    for issue in issues:
        key = (issue.code, issue.severity, issue.relative_path, issue.detail)
        if key in seen:
            continue
        seen.add(key)
        unique.append(issue)
    return unique
